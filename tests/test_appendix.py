"""The xlsx ingestion path: the cell map, the matching, and what it refuses.

Every workbook here is SYNTHETIC and every figure invented. The shapes are
the ones Pandora's appendix actually has -- a label repeated in two period
blocks, a label repeated inside one, a heading carrying the same text as
the line below it, a dash where a line does not apply, and a ratio on a
different sheet from its two components -- because those are what decide
whether label matching works.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
import yaml

from vss import manual as M
from vss.appendix import (
    AppendixError,
    CALENDAR_ENDS,
    extract,
    find_row,
    load_workbook,
    parse_header,
    parse_map,
    read_sheet,
    render_manual_yaml,
    run_appendix,
    verified_count,
    write_manual,
)

AS_OF = date(2026, 8, 24)
PERIODS = ["Q1 2026", "Q2 2026"]


def workbook(tmp_path: Path, sheets: dict, name="book.xlsx") -> Path:
    """Build an xlsx from {sheet name: [row, ...]} where a row is a list."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


#: One sheet with two period blocks, so the SAME label means two things.
FIN = [
    ["Consolidated income statement", *PERIODS],
    ["DKK million"],
    ["Revenue", 100, 110],
    ["Gross profit", 60, 66],
    ["Operating profit", 20, 22],
    ["Net profit for the period", 15, 16],
    [],
    ["Consolidated balance sheet", *PERIODS],
    ["Total assets", 500, 510],
    ["Property, plant and equipment", 80, 82],
    ["Cash", 30, 31],
    ["Total equity", 200, 205],
    ["Loans and borrowings", 150, 152],          # non-current
    ["Total non-current liabilities", 160, 162],
    ["Loans and borrowings", 40, 41],            # current
    [],
    ["Consolidated cash flow statement", *PERIODS],
    ["Operating profit", 20, 22],                # same label, other block
    ["Income taxes paid", -5, -6],
    ["Cash flows from operating activities, net", 25, 27],
    ["Actuarial gain/loss", "-", 3],             # a dash, not a zero
]
#: The ratio lives on a DIFFERENT sheet, as it does in the real appendix.
MARGINS = [
    ["Cost of sales and gross profit", *PERIODS],
    ["EBITDA"],                                  # a heading, no figures
    ["EBITDA", 30, 33],                          # the line
    ["EBIT margin", 0.2, 0.2],
]


def book(tmp_path: Path) -> Path:
    return workbook(tmp_path, {"Financial statements_appendix": FIN,
                               "Cost, GM, EBIT, EBITDA_appendix": MARGINS})


def cell_map(fields: dict, path=Path("map.yaml"), **overrides):
    document = {
        "ticker": "TEST.CO", "name": "Test A/S", "reporting_currency": "DKK",
        "quote_currency": "DKK", "sector": "Industrials",
        "reporting_frequency": "quarterly", "period_basis": "calendar",
        "document": "Test appendix", "fields": fields,
    }
    document.update(overrides)
    return parse_map(document, path=path)


FULL_FIELDS = {
    "revenue": {"sheet": "Financial statements_appendix",
                "section": "Consolidated income statement", "label": "Revenue"},
    "operating_income": {"sheet": "Financial statements_appendix",
                         "section": "Consolidated income statement",
                         "label": "Operating profit"},
    "op_margin": {"sheet": "Cost, GM, EBIT, EBITDA_appendix",
                  "label": "EBIT margin"},
}


# --- the header row --------------------------------------------------------


def test_period_headers_parse_into_the_schemas_own_labels():
    assert parse_header("Q2 2026") == "2026-Q2"
    assert parse_header("  Q1 2024  ") == "2026-Q1".replace("2026", "2024")
    assert parse_header("H1 2026") == "2026-H1"
    assert parse_header("FY 2025") == "2025-FY"


def test_a_header_that_is_not_a_period_is_not_a_period():
    for junk in ("Total", "2026", "Q5 2026", "", None, "Jan-Jun 2026"):
        assert parse_header(junk) is None


# --- matching on the label -------------------------------------------------


def test_a_label_the_sheet_does_not_have_stops_the_run(tmp_path):
    sheets = load_workbook(book(tmp_path))
    m = cell_map({"revenue": {"sheet": "Financial statements_appendix",
                              "label": "Total revenue"}})
    with pytest.raises(AppendixError) as exc:
        find_row(sheets["Financial statements_appendix"], m.fields[0])
    assert "no row carries that label" in str(exc.value)
    # and it names the near miss rather than silently taking it
    assert "'Revenue'" in str(exc.value)


def test_a_label_matching_twice_stops_the_run_and_names_both(tmp_path):
    """`Operating profit` is in the income statement AND the cash flow."""
    sheets = load_workbook(book(tmp_path))
    m = cell_map({"operating_income": {"sheet": "Financial statements_appendix",
                                       "label": "Operating profit"}})
    with pytest.raises(AppendixError) as exc:
        find_row(sheets["Financial statements_appendix"], m.fields[0])
    message = str(exc.value)
    assert "matches 2 rows" in message
    assert "Consolidated income statement" in message
    assert "Consolidated cash flow statement" in message
    assert "NOT resolved by taking the first" in message


def test_section_tells_two_blocks_apart(tmp_path):
    sheets = load_workbook(book(tmp_path))
    income = cell_map({"operating_income": {
        "sheet": "Financial statements_appendix",
        "section": "Consolidated income statement",
        "label": "Operating profit"}})
    row = find_row(sheets["Financial statements_appendix"], income.fields[0])
    assert row.section == "Consolidated income statement"


def test_after_tells_two_rows_inside_one_block_apart(tmp_path):
    """Both `Loans and borrowings` rows sit under the balance sheet."""
    sheets = load_workbook(book(tmp_path))
    current = cell_map({"financial_liabilities_current": {
        "sheet": "Financial statements_appendix",
        "label": "Loans and borrowings",
        "after": "Total non-current liabilities"}})
    row = find_row(sheets["Financial statements_appendix"], current.fields[0])
    assert row.values["2026-Q1"] == 40

    noncurrent = cell_map({"financial_liabilities_noncurrent": {
        "sheet": "Financial statements_appendix",
        "label": "Loans and borrowings", "after": "Total equity"}})
    assert find_row(sheets["Financial statements_appendix"],
                    noncurrent.fields[0]).values["2026-Q1"] == 150


def test_an_anchor_that_is_itself_ambiguous_is_refused(tmp_path):
    """Otherwise the ambiguity has moved rather than been resolved."""
    sheets = load_workbook(book(tmp_path))
    m = cell_map({"financial_liabilities_current": {
        "sheet": "Financial statements_appendix",
        "label": "Loans and borrowings", "after": "Operating profit"}})
    with pytest.raises(AppendixError, match="occurs 2 times"):
        find_row(sheets["Financial statements_appendix"], m.fields[0])


def test_a_heading_carrying_the_same_text_as_its_line_is_not_read(tmp_path):
    """The margin sheet has `EBITDA` as a block heading directly above
    `EBITDA` the line. A label row blank across every period is a heading."""
    sheets = load_workbook(book(tmp_path))
    m = cell_map({"ebitda": {"sheet": "Cost, GM, EBIT, EBITDA_appendix",
                             "label": "EBITDA"}})
    row = find_row(sheets["Cost, GM, EBIT, EBITDA_appendix"], m.fields[0])
    assert row.values["2026-Q1"] == 30


def test_nothing_matches_on_a_row_number(tmp_path):
    """Insert a row at the top and every figure must be unchanged."""
    shifted = [["Note: restated"], [], *FIN]
    a = load_workbook(book(tmp_path))["Financial statements_appendix"]
    b = load_workbook(workbook(
        tmp_path, {"Financial statements_appendix": shifted},
        name="shifted.xlsx"))["Financial statements_appendix"]
    spec = cell_map({"total_assets": {
        "sheet": "Financial statements_appendix",
        "section": "Consolidated balance sheet",
        "label": "Total assets"}}).fields[0]
    assert find_row(a, spec).values == find_row(b, spec).values


# --- what a cell is --------------------------------------------------------


def test_a_dash_is_data_missing_and_never_a_zero(tmp_path):
    sheets = load_workbook(book(tmp_path))
    rows = {r.label: r for r in sheets["Financial statements_appendix"].rows}
    m = cell_map({"revenue": {"sheet": "Financial statements_appendix",
                              "label": "Actuarial gain/loss"}})
    # The row exists and carries a figure in one period only.
    row = find_row(sheets["Financial statements_appendix"], m.fields[0])
    from vss.appendix import _numeric
    assert _numeric(row.values["2026-Q1"]) is None
    assert _numeric(row.values["2026-Q2"]) == 3.0


def test_a_cell_that_is_neither_a_number_nor_a_known_blank_stops_the_run():
    from vss.appendix import _numeric

    with pytest.raises(AppendixError, match="neither a number nor"):
        _numeric("c. 1,200")


# --- the map itself --------------------------------------------------------


def test_a_map_naming_a_field_the_schema_does_not_have_fails_at_load():
    with pytest.raises(AppendixError, match="the manual schema has no such field"):
        cell_map({"ebit_margin_adjusted": {"sheet": "s", "label": "l"}})


def test_a_map_field_without_a_label_fails():
    with pytest.raises(AppendixError, match="label is required"):
        cell_map({"revenue": {"sheet": "s"}})


def test_an_unknown_map_key_fails():
    with pytest.raises(AppendixError, match="unknown key"):
        cell_map({"revenue": {"sheet": "s", "label": "l"}}, isin="DK00")


def test_a_map_field_may_carry_a_note_and_it_is_kept():
    """The map is only a reviewable artefact if the close calls are in it."""
    m = cell_map({"cash_and_equivalents": {
        "sheet": "Financial statements_appendix", "label": "Cash",
        "note": "the balance sheet line, not the cash-flow one"}})
    assert "not the cash-flow one" in m.fields[0].note


def test_the_map_reports_every_schema_field_it_does_not_name():
    m = cell_map(FULL_FIELDS)
    assert "diluted_eps" in m.unmapped
    assert "lease_liabilities" in m.unmapped
    assert "revenue" not in m.unmapped


# --- periods ---------------------------------------------------------------


def test_periods_are_matched_across_sheets_by_label_not_by_position(tmp_path):
    """`op_margin` is on another sheet whose columns are in another order.

    Matching on the header text makes a misalignment impossible; matching
    on column position would silently pair Q1's margin with Q2's revenue.
    """
    reversed_margins = [
        ["Cost of sales and gross profit", "Q2 2026", "Q1 2026"],
        ["EBIT margin", 0.2, 0.2],
    ]
    path = workbook(tmp_path, {"Financial statements_appendix": FIN,
                               "Cost, GM, EBIT, EBITDA_appendix": reversed_margins})
    result = extract(cell_map(FULL_FIELDS), load_workbook(path))
    margin = result.by_field["op_margin"]
    assert margin.values == {"2026-Q1": 0.2, "2026-Q2": 0.2}


def test_period_end_is_derived_from_the_label_on_a_calendar_basis(tmp_path):
    text = render_manual_yaml(
        extract(cell_map(FULL_FIELDS), load_workbook(book(tmp_path))),
        workbook=Path("book.xlsx"), as_of=AS_OF)
    assert "period_end: 2026-06-30" in text
    assert "the appendix states no closing date" in text


def test_a_fiscal_basis_refuses_to_derive_a_period_end(tmp_path):
    """Where a fiscal period ends is a fact about the company (backlog B-3)."""
    m = cell_map(FULL_FIELDS, period_basis="fiscal")
    result = extract(m, load_workbook(book(tmp_path)))
    with pytest.raises(AppendixError, match="period_ends is missing"):
        render_manual_yaml(result, workbook=Path("b.xlsx"), as_of=AS_OF)


def test_a_fiscal_basis_uses_the_dates_the_map_states(tmp_path):
    m = cell_map(FULL_FIELDS, period_basis="fiscal",
                 period_ends={"2026-Q1": "2026-05-03", "2026-Q2": "2026-08-02"})
    text = render_manual_yaml(extract(m, load_workbook(book(tmp_path))),
                              workbook=Path("b.xlsx"), as_of=AS_OF)
    assert "period_end: 2026-05-03" in text and "period_end: 2026-08-02" in text


def test_a_column_of_the_wrong_kind_for_the_ticker_stops_the_run(tmp_path):
    half = [["Consolidated income statement", "H1 2026"], ["Revenue", 100]]
    path = workbook(tmp_path, {"Financial statements_appendix": half})
    m = cell_map({"revenue": {"sheet": "Financial statements_appendix",
                              "label": "Revenue"}})
    with pytest.raises(AppendixError, match="stores Q1/Q2/Q3/Q4 periods only"):
        extract(m, load_workbook(path))


# --- what comes out --------------------------------------------------------


def emitted(tmp_path):
    return render_manual_yaml(
        extract(cell_map(FULL_FIELDS), load_workbook(book(tmp_path))),
        workbook=Path("sources/book.xlsx"), as_of=AS_OF)


def test_every_figure_lands_unverified(tmp_path):
    text = emitted(tmp_path)
    assert "status: VERIFIED" not in text
    assert text.count("status: UNVERIFIED") == 6      # 3 fields x 2 periods


def test_the_page_carries_the_sheet_name_and_the_row_label(tmp_path):
    assert 'page: "Financial statements_appendix · Revenue"' in emitted(tmp_path)
    assert 'page: "Cost, GM, EBIT, EBITDA_appendix · EBIT margin"' in emitted(tmp_path)


def test_the_origin_is_nordic_xlsx_and_reaches_the_ranking_inputs(tmp_path):
    from vss import ranking

    text = emitted(tmp_path)
    assert "origin: nordic-xlsx" in text
    target = tmp_path / "TEST.CO.yaml"
    target.write_text(text, encoding="utf-8")
    parsed = M.parse_manual(yaml.safe_load(text), path=target)
    assert parsed.origin == "nordic-xlsx"
    record = M.as_record(parsed, as_of=AS_OF)
    assert record.origin == "nordic-xlsx"
    assert ranking.extract_inputs("TEST.CO", record).origin == "nordic-xlsx"


def test_the_output_is_byte_identical_on_a_re_run(tmp_path):
    """So a later appendix produces a diff a person can read."""
    assert emitted(tmp_path) == emitted(tmp_path)


def test_the_emitted_file_loads_through_the_manual_loader(tmp_path):
    text = emitted(tmp_path)
    parsed = M.parse_manual(yaml.safe_load(text), path=Path("TEST.CO.yaml"))
    assert [p.period for p in parsed.periods] == ["2026-Q1", "2026-Q2"]
    assert parsed.periods[-1].value("revenue") == 110


def test_a_figure_the_appendix_does_not_state_is_absent_never_derived(tmp_path):
    """`net_interest_paid` would be the sum of two stated rows. It is not
    summed: a figure derived from neighbours is a figure the sheet does not
    state, which is the one thing this path exists not to do."""
    result = extract(cell_map(FULL_FIELDS), load_workbook(book(tmp_path)))
    assert "net_interest_paid" in result.unmapped
    assert "net_interest_paid" not in emitted(tmp_path)


# --- validation before writing ---------------------------------------------


def test_a_unit_error_in_the_sheet_stops_the_emit(tmp_path):
    """A margin stated as a percentage would sail through cell reading and
    be caught by the manual loader's own unit contract."""
    percent = [["Cost of sales and gross profit", *PERIODS],
               ["EBIT margin", 20.0, 20.0]]
    path = workbook(tmp_path, {"Financial statements_appendix": FIN,
                               "Cost, GM, EBIT, EBITDA_appendix": percent})
    maps = tmp_path / "maps"; maps.mkdir()
    (maps / "TEST.CO.yaml").write_text(yaml.safe_dump({
        "ticker": "TEST.CO", "name": "Test A/S", "reporting_currency": "DKK",
        "quote_currency": "DKK", "sector": "Industrials",
        "fields": FULL_FIELDS}), encoding="utf-8")
    code, report = run_appendix(ticker="TEST.CO", workbook=path, write=True,
                               maps_dir=maps, manual_dir=tmp_path, as_of=AS_OF)
    assert code == 1
    assert "DOES NOT LOAD" in report and "FRACTION" in report
    assert not (tmp_path / "TEST.CO.yaml").exists(), "nothing may be written"


def test_a_cross_sheet_column_mismatch_shows_up_as_a_margin_gap(tmp_path):
    """The strongest check this path gets, and it is free.

    `op_margin` comes from a different sheet than its two components, so a
    column paired with the wrong quarter cannot reconcile.
    """
    wrong = [["Cost of sales and gross profit", *PERIODS],
             ["EBIT margin", 0.2, 0.9]]
    path = workbook(tmp_path, {"Financial statements_appendix": FIN,
                               "Cost, GM, EBIT, EBITDA_appendix": wrong})
    text = render_manual_yaml(extract(cell_map(FULL_FIELDS), load_workbook(path)),
                              workbook=Path("b.xlsx"), as_of=AS_OF)
    with pytest.raises(Exception, match="do not describe one measure"):
        M.parse_manual(yaml.safe_load(text), path=Path("TEST.CO.yaml"))


# --- writing ---------------------------------------------------------------


def test_writing_refuses_to_flatten_a_file_that_holds_verified_work(tmp_path):
    """The status flag records a reading only a person can do."""
    target = tmp_path / "TEST.CO.yaml"
    text = emitted(tmp_path).replace("status: UNVERIFIED", "status: VERIFIED", 2)
    target.write_text(text, encoding="utf-8")
    assert verified_count(target) == 2

    with pytest.raises(AppendixError) as exc:
        write_manual(emitted(tmp_path), target, as_of=AS_OF)
    assert "2 figure(s) already marked VERIFIED" in str(exc.value)
    assert "--force" in str(exc.value)
    assert "rule question and the owner's" in str(exc.value)
    assert target.read_text(encoding="utf-8") == text, "nothing may be touched"


def test_force_replaces_it_and_backs_it_up_first(tmp_path):
    target = tmp_path / "TEST.CO.yaml"
    target.write_text("old: file\n", encoding="utf-8")
    write_manual(emitted(tmp_path), target, as_of=AS_OF, force=True)
    backup = target.with_suffix(target.suffix + f".bak-{AS_OF.isoformat()}")
    assert backup.read_text(encoding="utf-8") == "old: file\n"
    assert "origin: nordic-xlsx" in target.read_text(encoding="utf-8")


def test_without_write_nothing_lands_on_disk_and_the_yaml_is_shown(tmp_path):
    maps = tmp_path / "maps"; maps.mkdir()
    (maps / "TEST.CO.yaml").write_text(yaml.safe_dump({
        "ticker": "TEST.CO", "name": "Test A/S", "reporting_currency": "DKK",
        "quote_currency": "DKK", "sector": "Industrials",
        "fields": FULL_FIELDS}), encoding="utf-8")
    code, report = run_appendix(ticker="TEST.CO", workbook=book(tmp_path),
                                maps_dir=maps, manual_dir=tmp_path, as_of=AS_OF)
    assert code == 0
    assert "NOT WRITTEN" in report
    assert "origin: nordic-xlsx" in report
    assert not (tmp_path / "TEST.CO.yaml").exists()


def test_the_emitted_file_carries_per_period_figures_and_no_sum(tmp_path):
    """E13's TTM sum lives in the ranking key and nowhere else. What the
    appendix writes is what the appendix states, quarter by quarter."""
    text = emitted(tmp_path)
    assert "value: 100" in text and "value: 110" in text      # the two quarters
    assert "value: 210" not in text                            # never their sum


# --- the units the run record divides on travel from the map to the file ----


def test_the_map_carries_the_unit_declarations_onto_the_written_file(tmp_path):
    m = cell_map(FULL_FIELDS, money_unit="millions", share_unit="whole")
    assert (m.money_unit, m.share_unit) == ("millions", "whole")
    text = render_manual_yaml(extract(m, load_workbook(book(tmp_path))),
                              workbook=Path("b.xlsx"), as_of=AS_OF)
    assert "money_unit: millions" in text and "share_unit: whole" in text
    # a map that does not say writes a file that does not say
    bare = render_manual_yaml(extract(cell_map(FULL_FIELDS), load_workbook(book(tmp_path))),
                              workbook=Path("b.xlsx"), as_of=AS_OF)
    assert "money_unit" not in bare and "share_unit" not in bare
    with pytest.raises(AppendixError, match="money_unit must be one of"):
        cell_map(FULL_FIELDS, money_unit="million")
